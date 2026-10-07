"""Deterministic table queries; no Python/SQL execution supplied by models."""
from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
import io
import json
from pathlib import Path

from .policy import ToolError


def load_table(path: Path, sheet: str | None):
    if path.suffix.lower() in {".csv", ".tsv"}:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("gb18030")
        rows = list(csv.reader(io.StringIO(text), delimiter="\t" if path.suffix.lower() == ".tsv" else ","))
    elif path.suffix.lower() == ".xlsx":
        import openpyxl
        with path.open("rb") as handle:
            book = openpyxl.load_workbook(handle, read_only=True, data_only=False)
            try:
                if sheet and sheet not in book.sheetnames:
                    raise ToolError("sheet_unavailable", "工作表不存在，可用工作表：" + ", ".join(book.sheetnames)[:1000])
                selected = book[sheet] if sheet else book[book.sheetnames[0]]
                rows = []
                for row in selected.iter_rows(values_only=True):
                    rows.append(list(row))
                sheet = selected.title
            finally:
                book.close()
    else:
        raise ToolError("unsupported_table", "query_table 支持 CSV、TSV、XLSX 原文件；PDF 表格请读取原文并检查对应页")
    if not rows:
        raise ToolError("empty_table", "表格为空")
    headers = [str(value or "").strip() for value in rows[0]]
    if len(set(headers)) != len(headers) or any(not h for h in headers):
        raise ToolError("ambiguous_table_header", "表头存在重复或空白，不能可靠按列计算")
    return headers, [{"_row": index, **{h: row[i] if i < len(row) else None for i, h in enumerate(headers)}}
                     for index, row in enumerate(rows[1:], start=2)], sheet


def number(value):
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
        if not parsed.is_finite():
            raise InvalidOperation
        return parsed
    except (InvalidOperation, ValueError):
        raise ToolError("non_numeric_cell", "计算列包含非数值或公式；请先用 select 检查原始单元格，不能将缺值当作零") from None


def query_table(path: Path, args: dict):
    headers, rows, sheet = load_table(path, args.get("sheet"))
    columns = args.get("columns") or headers
    value_column = args.get("value_column")
    group_by = args.get("group_by")
    requested = set(columns) | {f["column"] for f in args["filters"]} | {c for c in (value_column, group_by) if c}
    if not requested <= set(headers):
        raise ToolError("unknown_column", "列名无效，可用列：" + ", ".join(headers)[:1200])
    selected = rows
    for filt in args["filters"]:
        op, target, column = filt["op"], filt["value"], filt["column"]
        if op == "eq":
            selected = [r for r in selected if str(r[column]) == target]
        elif op == "contains":
            selected = [r for r in selected if target in str(r[column])]
        elif op == "gt":
            selected = [r for r in selected if number(r[column]) > number(target)]
        else:
            selected = [r for r in selected if number(r[column]) < number(target)]
    operation = args["operation"]
    if operation == "select":
        offset, limit = args["offset"], args["limit"]
        result = [{k: str(v) if v is not None else None for k, v in row.items() if k in columns or k == "_row"}
                  for row in selected[offset:offset + limit]]
        more = offset + limit < len(selected)
    else:
        if operation != "count" and not value_column:
            raise ToolError("missing_column", "数值计算必须指定 value_column")
        groups = {}
        for row in selected:
            groups.setdefault(str(row[group_by]) if group_by else "all", []).append(row)
        result, more = [], False
        if not groups and operation == "count":
            result.append({"group": "all", "value": "0", "rows": 0})
        for group, members in groups.items():
            values = [number(r[value_column]) for r in members] if operation != "count" else []
            value = (len(members) if operation == "count" else sum(values) if operation == "sum" else
                     sum(values) / len(values) if operation == "mean" else min(values) if operation == "min" else max(values))
            result.append({"group": group, "value": str(value), "rows": len(members)})
    return {"sheet": sheet, "columns": headers, "total_rows": len(rows), "matched_rows": len(selected),
            "operation": operation, "value_column": value_column, "group_by": group_by,
            "filters": args["filters"], "result": result, "truncated": more,
            "next_offset": args["offset"] + args["limit"] if more else None,
            "formula_policy": "公式作为原文返回，不执行、不使用可能过期的缓存数值"}
